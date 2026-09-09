from pathlib import Path
import openpyxl
import pytest
from backend.core import row,validate,parse

@pytest.mark.parametrize('raw,expected', [('35','35%'),('0.35','35%'),('35%','35%'),(None,'35%'),('350 บาท','35%')])
def test_discount_formats_reconcile(raw,expected):
    r=validate(row('x','x',1,'x','x','SKU-A',2,500,raw,gross=1000,net=650))
    assert r['status']=='ready'
    assert r['discount']==expected


def test_explicit_small_percentage_is_not_fraction():
    r=validate(row('x','x',1,'x','x','SKU-A',1,1000,'0.35%',net='996.50'))
    assert r['discount']=='0.35%' and r['status']=='ready'


def test_no_base_cannot_infer_missing_discount():
    r=validate(row('x','x',1,'x','x','SKU-A',1,100))
    assert r['status']=='ready' and r['discount']=='0%' and r['discount_defaulted']


def test_inconsistent_explicit_discount_is_blocked():
    r=validate(row('x','x',1,'x','x','SKU-A',1,100,'10%',net=70))
    assert r['status']=='blocked'


def test_unlabelled_fraction_without_net_defaults_zero():
    r=validate(row('x','x',1,'x','x','SKU-A',1,100,'0.35'))
    assert r['status']=='ready' and r['discount']=='0%' and r['discount_defaulted']


def test_outcast_accepts_percent_number_and_text(tmp_path):
    from backend.config import SAMPLES
    source=SAMPLES/'Sales Report Warrix - Aug 2026.xlsx'
    wb=openpyxl.load_workbook(source)
    cached=openpyxl.load_workbook(source,data_only=True)
    for i in [8,9,10]:
        for col in [10,11]:wb.active.cell(i,col,cached.active.cell(i,col).value)
    cached.close()
    for i,v in zip([8,9,10],[35,'35%',None]):wb.active.cell(i,6,v)
    # Assign None explicitly (cell(value=None) does not clear an existing cell).
    wb.active['F10']=None
    p=tmp_path/'outcast-variants.xlsx';wb.save(p);wb.close()
    rows,kind=parse(p,p.name)
    assert len(rows)==3 and all(r['discount']=='35%' for r in rows)
    assert all(not any('ยอดสุทธิ' in e for e in r['errors']) for r in rows)
