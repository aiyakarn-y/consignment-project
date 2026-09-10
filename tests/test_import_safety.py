import io,json
import openpyxl,pytest
from backend.core import parse_xlsx,row,validate
from backend.import_profiles import parse_profile,ImportProfile
from backend import app as service
from test_grouping import client,seed

def book(headers=('SKU','Qty','Price','Cost'),values=('A',2,100,60)):
    w=openpyxl.Workbook();w.active.title='Sales';w.active.append(headers);w.active.append(values)
    b=io.BytesIO();w.save(b);w.close();return b.getvalue()

def config(**changes):
    return dict(name='safe',sheet='Sales',header_row=1,customer='Shop',columns={'sku':'A','qty':'B','price':'C','cost_price':'D'},expected_headers={'A':'SKU','B':'Qty','C':'Price','D':'Cost'},**changes)

def test_header_drift_rejected():
    with pytest.raises(ValueError,match='หัวคอลัมน์'):
        parse_profile(book(('SKU','Price','Qty','Cost'),('A',100,2,60)),'drift.xlsx',config())

def test_duplicate_mapping_rejected():
    cfg=config();cfg['columns']['qty']='C'
    with pytest.raises(ValueError,match='ซ้ำ'):ImportProfile.model_validate(cfg)

def test_cost_is_not_automatically_settlement():
    rows,_=parse_profile(book(),'new.xlsx',config())
    assert rows[0]['net'] is None and rows[0]['discount']=='0%'
    assert rows[0]['calculated_net']=='200' and rows[0]['source_cost_total']=='120'

def test_explicit_cost_settlement():
    rows,_=parse_profile(book(),'new.xlsx',config(net_source='cost'))
    assert rows[0]['net']=='120' and rows[0]['discount']=='40%'

def multi(tmp_path):
    w=openpyxl.Workbook()
    for i in range(2):
        s=w.active if i==0 else w.create_sheet();s.title=f'Branch{i+1}'
        for key,val in {'C2':'REFERENCE','F2':'QTY','A4':1,'C4':f'SKU-{i+1}','D4':100,'E4':80,'F4':1,'G4':100,'I4':80}.items():s[key]=val
    w.create_sheet('Summary')['A1']='Summary'
    p=tmp_path/'multi.xlsx';w.save(p);w.close();return p

def test_multiple_sheets_require_choice(tmp_path):
    with pytest.raises(ValueError,match='เลือกชีต'):parse_xlsx(multi(tmp_path),'multi.xlsx')

def test_selected_sheets_all_read(tmp_path):
    rows,_=parse_xlsx(multi(tmp_path),'multi.xlsx',sheets=['Branch1','Branch2'])
    assert [r['sku'] for r in rows]==['SKU-1','SKU-2']
    assert [r['sheet'] for r in rows]==['Branch1','Branch2']

def test_master_replaces_fallback_but_preserves_manual_and_actual_sku(client):
    rows=[validate(row('x','s',1,'Shop','B','',1,100,partner='REF')),validate(row('x','s',2,'Shop','B','ACTUAL',1,100,partner='REF')),validate(row('x','s',3,'Shop','B','MANUAL',1,100,partner='REF'))]
    rows[2]['sku_origin']='ผู้ใช้กรอก'
    url=seed(client,rows)
    # Saving an unchanged fallback SKU while editing a discount must not promote it to a manual SKU.
    assert client.patch(url+'/rows/0',json={'sku':'REF','price':'100','qty':'1','discount':'20%'}).status_code==200
    result=client.post(url+'/master',files={'file':('master.csv',b'Customer,PartnerCode,SKU\nShop,REF,REAL\n')})
    assert result.json()['resolved']==1
    stored=service.load(url.split('/')[-1])['rows']
    assert [r['sku'] for r in stored]==['REAL','ACTUAL','MANUAL']
    assert stored[0]['edits'][-1]['before']['sku']=='REF'
    new={'rows':[validate(row('x','s',1,'Shop','B','',1,100,partner='REF'))]}
    assert service.mapped(new)['rows'][0]['sku']=='REAL'


def test_preview_binds_headers_and_legacy_profiles_require_review(client):
    cfg=config();cfg.pop('expected_headers')
    assert client.post('/api/profiles',json={'config':cfg}).status_code==400
    preview=client.post('/api/profiles/preview',files={'file':('new.xlsx',book())},data={'config':json.dumps(cfg)})
    assert preview.status_code==200
    bound=preview.json()['config']
    assert bound['expected_headers']==config()['expected_headers'] and bound['net_source']=='mapped'
    saved=client.post('/api/profiles',json={'config':bound}).json()
    b=client.post('/api/batches',json={}).json();url='/api/batches/'+b['id']
    changed=book(('SKU','Price','Qty','Cost'),('A',100,2,60))
    result=client.post(url+'/upload',files={'files':('changed.xlsx',changed)},data={'profile_id':saved['id']}).json()
    assert result['rows']==0 and result['file_errors']==1
    assert 'หัวคอลัมน์' in result['files'][0]['error']
    assert client.get(url+'/export').status_code==409


def test_api_multisheet_inventory_selection_and_skip_report(client,tmp_path):
    data=multi(tmp_path).read_bytes();b=client.post('/api/batches',json={}).json();url='/api/batches/'+b['id']
    checks=client.post(url+'/check-upload',files={'files':('multi.xlsx',data)}).json()
    assert [s['supported'] for s in checks[0]['sheets']]==[True,True,False]
    result=client.post(url+'/upload',files={'files':('multi.xlsx',data)},data={'sheet_selection':json.dumps({'0':['Branch1','Branch2']})}).json()
    assert result['rows']==2 and result['ready']==2
    review=result['files'][0]['sheet_review']
    assert [s['rows'] for s in review]==[1,1,0]
    assert [s['gross'] for s in review]==['100','100','0']
    assert review[2]['selected'] is False
    assert client.get(url+'/export').status_code==200


def test_profile_multiple_sheets_checks_every_header():
    w=openpyxl.load_workbook(io.BytesIO(book()));s=w.copy_worksheet(w.active);s.title='Other';s['B1']='Price';s['C1']='Qty'
    b=io.BytesIO();w.save(b);w.close()
    with pytest.raises(ValueError,match='หัวคอลัมน์ชีต Other'):
        parse_profile(b.getvalue(),'multi.xlsx',config(),sheets=['Sales','Other'])


@pytest.mark.parametrize('sheets',[[],['Branch1','Branch1'],['Summary'],['missing']])
def test_reject_invalid_sheet_selection(tmp_path,sheets):
    with pytest.raises(ValueError):parse_xlsx(multi(tmp_path),'multi.xlsx',sheets=sheets)


def test_missing_cost_when_explicitly_selected_is_rejected():
    with pytest.raises(ValueError,match='ไม่มีต้นทุน'):
        parse_profile(book(values=('A',2,100,None)),'empty-cost.xlsx',config(net_source='cost'))


def test_master_preview_does_not_write(client):
    url=seed(client,[validate(row('x','s',1,'Shop','B','',1,100,partner='REF'))])
    before=service.load(url.split('/')[-1])
    result=client.post(url+'/master',files={'file':('master.csv',b'Customer,PartnerCode,SKU\nShop,REF,REAL\n')},data={'preview':'true'}).json()
    assert result['resolved']==1 and result['preview'] is True
    assert service.load(url.split('/')[-1])==before
    with service.db() as c:assert c.execute('SELECT COUNT(*) FROM mappings').fetchone()[0]==0
