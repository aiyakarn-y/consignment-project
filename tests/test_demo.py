from pathlib import Path
from decimal import Decimal as D
from copy import deepcopy
import hashlib, io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from backend import app as service
from backend.core import parse, row, validate, percent, reverse_discount, export_xlsx, money

from backend.config import SAMPLES as BASE, TEMPLATE

@pytest.fixture(scope='module')
def reports():
    return {kind:rows for p in BASE.iterdir() if p.suffix.lower() in ('.xlsx','.pdf') and p.name!='Consign_sample.xlsx' and not p.name.startswith('~$') for rows,kind in [parse(p,p.name)]}

@pytest.mark.parametrize('kind,count,qty,gross,net',[
    ('Big C',3770,'4606','1069864','802529.70'),
    ('Outlet G624',70,'91','32196',None),
    ('KING POWER',43,'49','48207','29284.54'),
    ('Outcast',3,'3','2570','1670.50'),
    ('PDF ส่งเงินคืน',10,'46','39840','35856'),
])
def test_fixture_totals(reports,kind,count,qty,gross,net):
    rows=reports[kind]
    assert len(rows)==count
    assert sum(D(r['qty']) for r in rows)==D(qty)
    assert sum(D(r['gross']) for r in rows)==D(gross)
    if net:assert sum(D(r['net']) for r in rows)==D(net)


def test_preserve_dates_returns_branches(reports):
    rows=reports['Big C']
    assert len({r['branch_code'] for r in rows})==40
    assert sum(D(r['qty'])<0 for r in rows)==7
    assert sum(r['date'].endswith('07-2026') for r in rows)==1339
    branch=[r for r in rows if r['branch_code']=='11107']
    assert len(branch)==481
    assert sum(D(r['qty']) for r in branch)==551


def test_source_references_are_exportable(reports):
    assert all(r['sku'] and r['status']=='ready' for rows in reports.values() for r in rows)
    assert all(r['discount']=='0%' and r['discount_defaulted'] for r in reports['Outlet G624'])
    assert next(r for r in reports['Big C'] if r['line']==3241)['sku']=='103192895'
    assert {r['sku'] for r in reports['Outcast']}=={'26348','31203','26353'}


def test_reverse_discount_reconciles_every_king_power_row(reports):
    for r in reports['KING POWER']:
        assert r['discount']!='39.25%'
        assert r['discount_origin']=='คำนวณย้อนกลับ'
        assert money(D(r['price'])*D(r['qty'])*percent(r['discount']))==money(D(r['net']))
        assert r['status']=='ready'


def test_pdf_reverse_exact(reports):
    assert all(r['discount']=='10%' for r in reports['PDF ส่งเงินคืน'])


def test_compound_discount_and_return():
    assert percent('20%+5%')==D('.76')
    r=row('x','x',1,'x','x','SKU-A',-2,100,'20%+5%',net=-152)
    assert validate(r)['status']=='ready'
    assert r['calculated_net']=='-152'

@pytest.mark.parametrize('price,qty,discount,net',[(0,1,'0%',None),(10,0,'0%',0),(10,1,'25%',8),(10,1,'NaN',5),(10,1,'101%',0)])
def test_invalid_or_unreconciled_rows_block(price,qty,discount,net):
    r=dict(sku='SKU',price=str(price),qty=str(qty),discount=discount,net=str(net) if net is not None else None)
    assert validate(r)['status']=='blocked'


def test_export_preserves_template_and_reconstructs_net(tmp_path,reports):
    template=TEMPLATE;before=hashlib.sha256(template.read_bytes()).hexdigest()
    rows=deepcopy(reports['KING POWER']+reports['PDF ส่งเงินคืน'])
    path=tmp_path/'result.xlsx';export_xlsx(rows,template,path)
    wb=openpyxl.load_workbook(path);src=openpyxl.load_workbook(template);ws=wb['Sample']
    assert wb.sheetnames==['Sample']
    assert ws.max_column==4 and ws.max_row==54
    assert [c.value for c in ws[1]]==['Items','Price','Qty','Discount']
    for i,r in enumerate(rows,2):
        sku,p,q,discount=[c.value for c in ws[i]]
        assert sku==r['sku']
        expected_discount='+'.join(f'{money(D(part.rstrip(chr(37)))):.2f}%' for part in r['discount'].split('+'))
        assert discount==expected_discount
        for col in range(1,5): assert ws.cell(i,col)._style==src['Sample'].cell(2,col)._style
    for col in 'ABCD':assert ws.column_dimensions[col].width==src['Sample'].column_dimensions[col].width
    assert hashlib.sha256(template.read_bytes()).hexdigest()==before
    wb.close();src.close()


def test_blocked_export(tmp_path,reports):
    invalid=deepcopy(reports['Outlet G624']);invalid[0].update(sku='',partner='',raw={})
    with pytest.raises(ValueError):export_xlsx(invalid,TEMPLATE,tmp_path/'invalid.xlsx')

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(service,'DATA',tmp_path)
    monkeypatch.setattr(service,'DB',tmp_path/'test.db')
    return TestClient(service.app)


def test_api_full_sample_review_export(client):
    b=client.post('/api/batches',json={'period':'2026-08'}).json();url='/api/batches/'+b['id']
    res=client.post(url+'/sample');assert res.status_code==200
    summary=res.json();assert summary['rows']==3896 and summary['blocked']==0
    assert client.get(url+'/export').status_code==200
    assert client.post(url+'/sample').status_code==409
    res=client.post(url+'/discount-rule',json={'customer':'Outlet G624','discount':'20%+5%','remember':True})
    assert res.status_code==200 and res.json()['updated']==70
    final=client.get(url).json();assert final['blocked']==0 and final['exportable']
    res=client.get(url+'/export');assert res.status_code==200
    wb=openpyxl.load_workbook(io.BytesIO(res.content));values=list(wb['Sample'].iter_rows(min_row=2,values_only=True))
    assert len(values)==final['merged_rows']<3896
    assert len({r[0] for r in values})==len(values)
    assert sum(D(str(r[2])) for r in values)==4795
    wb.close()
    # Profiles work on the next report batch; test values never touch the demo DB.
    b2=client.post('/api/batches',json={'period':'2026-08'}).json()
    assert client.post('/api/batches/'+b2['id']+'/sample').json()['blocked']==0


def test_api_net_edit_reverse_and_audit(client):
    b=client.post('/api/batches',json={}).json();url='/api/batches/'+b['id']
    client.post(url+'/sample')
    r=next(r for r in service.load(b['id'])['rows'] if r['customer']=='Outlet G624' and r['sku'])
    net=str(D(r['gross'])*D('.9'))
    result=client.patch(url+'/rows/'+r['id'],json={'net':net}).json()
    assert result['discount']=='10%' and result['status']=='ready'
    assert result['edits'][-1]['after']['net']==net
    assert client.patch(url+'/rows/'+r['id'],json={'price':'NaN'}).status_code==400


def test_master_conflict_is_atomic(client):
    b=client.post('/api/batches',json={}).json();url='/api/batches/'+b['id']
    conflict=b'Barcode,SKU\n123,SKU-A\n123,SKU-B\n'
    assert client.post(url+'/master',files={'file':('m.csv',conflict)}).status_code==400
    with service.db() as c:assert c.execute('SELECT count(*) FROM mappings').fetchone()[0]==0


def test_bad_files_and_failed_source_blocks_export(client,tmp_path):
    b=client.post('/api/batches',json={}).json();url='/api/batches/'+b['id']
    assert client.post(url+'/upload',files={'files':('bad.xlsx',b'wrong')}).status_code==400
    w=openpyxl.Workbook();w.active['A1']='Unknown';path=tmp_path/'unknown.xlsx';w.save(path)
    result=client.post(url+'/upload',files={'files':('unknown.xlsx',path.read_bytes())})
    assert result.status_code==200 and result.json()['file_errors']==1
    assert client.get(url+'/export').status_code==409
    f=result.json()['files'][0]
    assert client.delete(url+'/files/'+f['id']).json()['file_errors']==0


def test_ready_only_export_excludes_blocked_and_preserves_batch(client):
    b=client.post('/api/batches',json={}).json();url='/api/batches/'+b['id']
    client.post(url+'/sample')
    before=service.load(b['id'])
    expected=[r for r in before['rows'] if validate(r)['status']=='ready']
    response=client.get(url+'/export?scope=ready')
    assert response.status_code==200,response.text
    assert response.headers['x-exported-rows']=='1008'
    assert response.headers['x-exported-source-rows']=='3896'
    assert response.headers['x-skipped-rows']=='0'
    assert response.headers['x-skipped-files']=='0'
    assert '_ready.xlsx' in response.headers['content-disposition']
    wb=openpyxl.load_workbook(io.BytesIO(response.content));ws=wb['Sample']
    assert wb.sheetnames==['Sample'] and ws.max_column==4 and ws.max_row==1009
    values=list(ws.iter_rows(min_row=2,values_only=True))
    assert len({r[0] for r in values})==len(values)
    by_sku={}
    for r in expected:
        totals=by_sku.setdefault(r['sku'],[D(0),D(0),D(0)])
        totals[0]+=D(r['qty']);totals[1]+=money(D(r['gross']));totals[2]+=money(D(r['calculated_net']))
    rounding_delta=D(0)
    for sku,p,q,d in values:
        target=by_sku[sku]
        assert D(str(q))==target[0]
        assert money(D(str(p))*D(str(q)))==target[1]
        assert all(len(part.rstrip('%').split('.')[1])==2 for part in d.split('+'))
        rounding_delta+=money(D(str(p))*D(str(q))*percent(d))-target[2]
    assert rounding_delta==D('1.41')
    assert D(response.headers['X-Export-Rounding-Delta'])==rounding_delta
    preview=client.get(url+'/merged?size=100').json()
    assert preview['total']==1008 and len(preview['rows'])==100
    wb.close()
    assert service.load(b['id'])==before
    assert client.get(url+'/export').status_code==200
    assert client.get(url+'/export?scope=invalid').status_code==422


def test_ready_export_empty_or_failed_file_reporting(client,tmp_path):
    b=client.post('/api/batches',json={}).json();url='/api/batches/'+b['id']
    assert client.get(url+'/export?scope=ready').status_code==409
    outlet=next(BASE.glob('G624*.xlsx'))
    client.post(url+'/upload',files={'files':(outlet.name,outlet.read_bytes())})
    assert client.get(url+'/export?scope=ready').status_code==200
    king=BASE/'SALES WIX AUG 26.xlsx'
    client.post(url+'/upload',files={'files':(king.name,king.read_bytes())})
    w=openpyxl.Workbook();w.active['A1']='unknown';buf=io.BytesIO();w.save(buf)
    client.post(url+'/upload',files={'files':('unknown.xlsx',buf.getvalue())})
    response=client.get(url+'/export?scope=ready')
    assert response.status_code==200
    assert int(response.headers['x-exported-rows'])==client.get(url).json()['merged_rows']
    assert response.headers['x-exported-source-rows']=='113'
    assert response.headers['x-skipped-rows']=='0'
    assert response.headers['x-skipped-files']=='1'
    assert client.get(url+'/export').status_code==409
