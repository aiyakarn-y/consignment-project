"""Explicit tabular Excel import profiles; never evaluate spreadsheet formulas."""
import io,re,zipfile
from decimal import Decimal
from typing import Literal
from pydantic import BaseModel,Field,ConfigDict,model_validator
import openpyxl
from backend.core import row,validate,dec,fmt,code,discount_text,reverse_discount

FIELDS={'sku','partner','barcode','qty','price','gross','net','cost_price','cost_total','discount','mg','customer','branch','branch_code','date'}

class Valuation(BaseModel):
    model_config=ConfigDict(extra='forbid')
    sale_basis:Literal['unknown','included','excluded']='unknown'
    cost_basis:Literal['unknown','included','excluded']='unknown'
    settlement_basis:Literal['unknown','included','excluded']='unknown'
    vat_rate:str='7'
    note:str=Field(default='',max_length=500)
    @model_validator(mode='after')
    def check(self):
        n=dec(self.vat_rate)
        if n is None or not 0<=n<=100:raise ValueError('อัตรา VAT ต้องอยู่ระหว่าง 0–100')
        self.vat_rate=fmt(n)
        return self

class ImportProfile(BaseModel):
    model_config=ConfigDict(extra='forbid')
    name:str=Field(min_length=1,max_length=100)
    sheet:str=Field(min_length=1,max_length=100)
    header_row:int=Field(default=1,ge=1,le=1000)
    end_row:int|None=Field(default=None,ge=2,le=50000)
    columns:dict[str,str]
    customer:str=Field(default='',max_length=150)
    branch:str=Field(default='',max_length=150)
    discount_mode:Literal['auto','percent','fraction','amount']='auto'
    valuation:Valuation=Field(default_factory=Valuation)
    expected_headers:dict[str,str]=Field(default_factory=dict)
    net_source:Literal['mapped','cost','none']='mapped'
    @model_validator(mode='after')
    def check(self):
        self.name=self.name.strip();self.customer=self.customer.strip()
        if not self.name:raise ValueError('ระบุชื่อรูปแบบ')
        self.columns={k:v.strip().upper() for k,v in self.columns.items() if v.strip()}
        if len(set(self.columns.values()))!=len(self.columns):raise ValueError('คอลัมน์ Mapping ซ้ำ กรุณาเลือกคอลัมน์แยกให้ถูกต้อง')
        if self.net_source=='cost' and not {'cost_price','cost_total'}&self.columns.keys():raise ValueError('ต้องระบุคอลัมน์ต้นทุนเมื่อเลือกใช้ต้นทุนเป็นยอดส่งคืน')
        if set(self.columns)-FIELDS:raise ValueError('มีช่อง Mapping ที่ไม่รองรับ')
        if any(not re.fullmatch('[A-Z]{1,3}',v) or openpyxl.utils.column_index_from_string(v)>200 for v in self.columns.values()):raise ValueError('คอลัมน์ต้องอยู่ระหว่าง A–GR')
        if not {'sku','partner'}&self.columns.keys():raise ValueError('ต้องมี SKU หรือรหัสสินค้าอ้างอิง')
        if 'qty' not in self.columns or not {'price','gross'}&self.columns.keys():raise ValueError('ต้องมีจำนวน และราคาขายต่อหน่วยหรือยอดขายรวม')
        if not self.customer and 'customer' not in self.columns:raise ValueError('ระบุลูกค้า หรือเลือกคอลัมน์ลูกค้า')
        if self.end_row and self.end_row<=self.header_row:raise ValueError('แถวสุดท้ายต้องอยู่หลังหัวตาราง')
        return self

def workbook(content):
    if len(content)>25*1024*1024:raise ValueError('Excel ต้องไม่เกิน 25 MB')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(f.file_size for f in z.infolist())>150*1024*1024:raise ValueError('Excel หลังคลายข้อมูลใหญ่เกิน 150 MB')
        return openpyxl.load_workbook(io.BytesIO(content),data_only=True,read_only=True)
    except (zipfile.BadZipFile,KeyError):raise ValueError('ไฟล์ Excel ไม่สมบูรณ์')

def inspect_workbook(content,sheet='',header_row=1):
    w=workbook(content)
    try:
        if not 1<=header_row<=1000:raise ValueError('หัวตารางต้องเป็นแถว 1–1000')
        if sheet and sheet not in w.sheetnames:raise ValueError('ไม่พบชีต')
        s=w[sheet] if sheet else w.active
        cols=min(s.max_column or 1,200)
        rows=list(s.iter_rows(min_row=header_row,max_row=header_row+5,max_col=cols,values_only=True))
        return dict(sheets=[dict(name=x.title,rows=x.max_row,columns=x.max_column) for x in w],sheet=s.title,header_row=header_row,columns=[dict(column=openpyxl.utils.get_column_letter(i+1),label=str(v or '')) for i,v in enumerate(rows[0])],sample=[[str(v) if v is not None else '' for v in line] for line in rows[1:]])
    finally:w.close()

def header_text(value):
    return ' '.join(str(value or '').split())


def require_profile_headers(cfg):
    if set(cfg.expected_headers)!=set(cfg.columns.values()) or any(not header_text(v) for v in cfg.expected_headers.values()):
        raise ValueError('ยังไม่ได้ยืนยันหัวคอลัมน์ กรุณาดูตัวอย่าง Mapping และบันทึกรูปแบบใหม่ก่อนนำเข้า')


def check_headers(sheet,cfg):
    require_profile_headers(cfg)
    for column,expected in cfg.expected_headers.items():
        actual=header_text(sheet[f'{column}{cfg.header_row}'].value)
        if actual!=header_text(expected):
            raise ValueError(f'หัวคอลัมน์ชีต {sheet.title} ช่อง {column} เปลี่ยน: คาดว่า {expected} แต่พบ {actual or "(ว่าง)"} กรุณาตรวจ Mapping ใหม่')


def bind_profile_headers(content,profile):
    cfg=ImportProfile.model_validate(profile)
    w=workbook(content)
    try:
        if cfg.sheet not in w.sheetnames:raise ValueError('ไม่พบชีตที่ตั้งค่า: '+cfg.sheet)
        cfg.expected_headers={c:header_text(w[cfg.sheet][f'{c}{cfg.header_row}'].value) for c in cfg.columns.values()}
        require_profile_headers(cfg)
        return cfg.model_dump()
    finally:w.close()


def inspect_sheet_choices(content,profile=None):
    from backend.core import sheet_kind
    w=workbook(content)
    try:
        result=[]
        cfg=ImportProfile.model_validate(profile) if profile else None
        for sheet in w:
            kind=sheet_kind(sheet) if cfg is None else 'รูปแบบ: '+cfg.name
            reason=''
            if cfg:
                try:check_headers(sheet,cfg)
                except ValueError as ex:kind='';reason=str(ex)
            result.append(dict(name=sheet.title,kind=kind,supported=bool(kind),reason=reason or ('' if kind else 'ชีตประกอบ/ไม่รองรับรูปแบบนี้'),suggested=bool(kind) and (cfg is None or sheet.title==cfg.sheet)))
        return result
    finally:w.close()


def parse_profile(content,name,profile,sheets=None):
    cfg=ImportProfile.model_validate(profile);w=workbook(content)
    try:
        selected=[cfg.sheet] if sheets is None else sheets
        if not selected or len(selected)!=len(set(selected)) or any(name not in w.sheetnames for name in selected):raise ValueError('เลือกชีตที่มีอยู่จริงและไม่ซ้ำอย่างน้อยหนึ่งชีต')
        out=[]
        for sheet_name in selected:
            s=w[sheet_name]
            first_row=len(out)
            check_headers(s,cfg)
            if (s.max_row or 0)>50000 or (s.max_column or 0)>200:raise ValueError('รูปแบบตารางรองรับไม่เกิน 50,000 แถว / 200 คอลัมน์')
            end=cfg.end_row or s.max_row
            columns={k:openpyxl.utils.column_index_from_string(v)-1 for k,v in cfg.columns.items()}
            for i,cells in enumerate(s.iter_rows(min_row=cfg.header_row+1,max_row=end,values_only=True),cfg.header_row+1):
                values={k:cells[j] if j<len(cells) else None for k,j in columns.items()}
                if all(v is None or str(v).strip()=='' for v in values.values()):continue
                raw={openpyxl.utils.get_column_letter(j+1):str(v) for j,v in enumerate(cells) if v is not None}
                try:
                    q=dec(values.get('qty'));cost_unit=dec(values.get('cost_price'));cost_total=dec(values.get('cost_total'))
                    if cost_total is None and cost_unit is not None and q is not None:cost_total=cost_unit*q
                    net=dec(values.get('net')) if cfg.net_source=='mapped' else None
                    if cfg.net_source=='cost':
                        if cost_total is None:raise ValueError('ไม่มีต้นทุนสำหรับยอดส่งคืนที่เลือก กรุณาตรวจต้นทาง')
                        net=cost_total
                    discount=values.get('discount')
                    if discount is not None and str(discount).strip():
                        if cfg.discount_mode=='fraction':discount=fmt(dec(discount)*100)+'%'
                        elif cfg.discount_mode=='percent':discount=discount_text(discount)
                        elif cfg.discount_mode=='amount':
                            gross=dec(values.get('gross'))
                            if gross is None and q is not None and dec(values.get('price')) is not None:gross=q*dec(values['price'])
                            if gross is None or gross==0:raise ValueError('ส่วนลดจำนวนเงินต้องมียอดขายรวมไม่เป็นศูนย์')
                            discount=reverse_discount(gross,gross-dec(discount))
                    r=row(name,s.title,i,code(values.get('customer')) or cfg.customer,code(values.get('branch')) or cfg.branch,code(values.get('sku')),q,values.get('price'),discount,values.get('gross'),net,values.get('barcode'),values.get('partner'),date=values.get('date'),raw=raw)
                    r.update(branch_code=code(values.get('branch_code')),valuation=cfg.valuation.model_dump(),custom_source=dict(cost_price=fmt(cost_unit),cost_total=fmt(cost_total),qty=fmt(q),gp=r['discount'] if discount not in (None,'') else None,mg=code(values.get('mg')) or None),import_profile=cfg.name,net_source=cfg.net_source)
                    out.append(validate(r))
                except (ValueError,TypeError,ArithmeticError) as ex:raise ValueError(f'แถว {i}: {ex}')
            if len(out)==first_row:raise ValueError('ไม่พบข้อมูลหลังแถวหัวตารางในชีตที่เลือก: '+sheet_name)
        if not out:raise ValueError('ไม่พบข้อมูลหลังแถวหัวตาราง')
        return out,'รูปแบบ: '+cfg.name
    finally:w.close()
